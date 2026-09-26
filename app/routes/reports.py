import io
import csv
from datetime import datetime

from flask import Blueprint, send_file
from flask_login import login_required, current_user

from app.models import Task
from app import db

# ⚠️ Nome do blueprint precisa ser EXATAMENTE 'reports'
reports_bp = Blueprint('reports', __name__, url_prefix='/reports')


# ============================================================
# HELPERS
# ============================================================

def _get_tasks_and_stats():
    tasks = Task.query.order_by(Task.created_at.desc()).all()

    total = len(tasks)
    by_status = {
        'A Fazer':   len([t for t in tasks if t.status == 'todo']),
        'Fazendo':   len([t for t in tasks if t.status == 'doing']),
        'Concluído': len([t for t in tasks if t.status == 'done']),
        'Ajustes':   len([t for t in tasks if t.status == 'adjust']),
    }
    by_priority = {
        'Alta':  len([t for t in tasks if t.priority == 'alta']),
        'Média': len([t for t in tasks if t.priority == 'media']),
        'Baixa': len([t for t in tasks if t.priority == 'baixa']),
    }
    grouped = {}
    for t in tasks:
        key = t.assignee.username if t.assignee else 'Sem responsável'
        if key not in grouped:
            grouped[key] = {'total': 0, 'todo': 0, 'doing': 0, 'done': 0, 'adjust': 0}
        grouped[key]['total'] += 1
        grouped[key][t.status] += 1

    return tasks, total, by_status, by_priority, grouped


def _filename(ext):
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    return f'relatorio_kanban_{ts}.{ext}'


STATUS_LABEL = {
    'todo': 'A Fazer',
    'doing': 'Fazendo',
    'done': 'Concluído',
    'adjust': 'Ajustes',
}


# ============================================================
# CSV
# ============================================================

@reports_bp.route('/csv')
@login_required
def export_csv():
    tasks, total, by_status, by_priority, grouped = _get_tasks_and_stats()

    output = io.StringIO()
    output.write('\ufeff')  # BOM para Excel PT-BR

    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_MINIMAL)

    writer.writerow(['RELATÓRIO QUADRO DE TAREFAS'])
    writer.writerow(['Gerado em', datetime.now().strftime('%d/%m/%Y %H:%M:%S')])
    writer.writerow(['Gerado por', current_user.username])
    writer.writerow([])

    writer.writerow(['RESUMO POR STATUS'])
    writer.writerow(['Status', 'Quantidade'])
    for k, v in by_status.items():
        writer.writerow([k, v])
    writer.writerow(['TOTAL', total])
    writer.writerow([])

    writer.writerow(['RESUMO POR PRIORIDADE'])
    writer.writerow(['Prioridade', 'Quantidade'])
    for k, v in by_priority.items():
        writer.writerow([k, v])
    writer.writerow([])

    writer.writerow(['POR RESPONSÁVEL'])
    writer.writerow(['Responsável', 'Total', 'A Fazer', 'Fazendo', 'Concluído', 'Ajustes'])
    for name, data in grouped.items():
        writer.writerow([name, data['total'], data['todo'],
                         data['doing'], data['done'], data['adjust']])
    writer.writerow([])

    writer.writerow(['TAREFAS'])
    writer.writerow(['ID', 'Título', 'Descrição', 'Status', 'Prioridade',
                     'Criado por', 'Responsável', 'Criada em', 'Concluída em'])
    for t in tasks:
        writer.writerow([
            t.id,
            t.title,
            (t.description or '').replace('\n', ' ').replace('\r', ' '),
            STATUS_LABEL.get(t.status, t.status),
            t.priority.capitalize(),
            t.owner.username if t.owner else '—',
            t.assignee.username if t.assignee else '—',
            t.created_at.strftime('%d/%m/%Y %H:%M'),
            t.completed_at.strftime('%d/%m/%Y %H:%M') if t.completed_at else '',
        ])

    output.seek(0)
    data = io.BytesIO(output.getvalue().encode('utf-8'))

    return send_file(
        data,
        mimetype='text/csv; charset=utf-8',
        as_attachment=True,
        download_name=_filename('csv'),
    )


# ============================================================
# EXCEL
# ============================================================

@reports_bp.route('/xlsx')
@login_required
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    tasks, total, by_status, by_priority, grouped = _get_tasks_and_stats()

    wb = Workbook()

    header_font = Font(bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill('solid', fgColor='343A40')
    title_font = Font(bold=True, size=14, color='0D6EFD')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin = Side(border_style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(ws, row, cols):
        for col in range(1, cols + 1):
            c = ws.cell(row=row, column=col)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center
            c.border = border

    def autosize(ws):
        for col in ws.columns:
            max_len = 10
            letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 2, 50)

    # ----- Aba 1: Resumo -----
    ws = wb.active
    ws.title = 'Resumo'

    ws['A1'] = 'RELATÓRIO QUADRO DE TAREFAS'
    ws['A1'].font = title_font
    ws.merge_cells('A1:B1')
    ws['A2'] = 'Gerado em'
    ws['B2'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
    ws['A3'] = 'Gerado por'
    ws['B3'] = current_user.username

    ws['A5'] = 'RESUMO POR STATUS'
    ws['A5'].font = Font(bold=True, size=12)
    ws['A6'] = 'Status'
    ws['B6'] = 'Quantidade'
    style_header(ws, 6, 2)
    row = 7
    for k, v in by_status.items():
        ws.cell(row=row, column=1, value=k).border = border
        ws.cell(row=row, column=2, value=v).border = border
        row += 1
    ws.cell(row=row, column=1, value='TOTAL').font = Font(bold=True)
    ws.cell(row=row, column=2, value=total).font = Font(bold=True)

    row += 2
    ws.cell(row=row, column=1, value='RESUMO POR PRIORIDADE').font = Font(bold=True, size=12)
    row += 1
    ws.cell(row=row, column=1, value='Prioridade')
    ws.cell(row=row, column=2, value='Quantidade')
    style_header(ws, row, 2)
    row += 1
    for k, v in by_priority.items():
        ws.cell(row=row, column=1, value=k).border = border
        ws.cell(row=row, column=2, value=v).border = border
        row += 1
    autosize(ws)

    # ----- Aba 2: Por Responsável -----
    ws2 = wb.create_sheet('Por Responsável')
    headers = ['Responsável', 'Total', 'A Fazer', 'Fazendo', 'Concluído', 'Ajustes']
    ws2.append(headers)
    style_header(ws2, 1, len(headers))
    for name, data in grouped.items():
        ws2.append([name, data['total'], data['todo'],
                    data['doing'], data['done'], data['adjust']])
    for r in ws2.iter_rows(min_row=2):
        for c in r:
            c.border = border
            c.alignment = center
        r[0].alignment = left
    autosize(ws2)

    # ----- Aba 3: Tarefas -----
    ws3 = wb.create_sheet('Tarefas')
    headers = ['ID', 'Título', 'Descrição', 'Status', 'Prioridade',
               'Criado por', 'Responsável', 'Criada em', 'Concluída em']
    ws3.append(headers)
    style_header(ws3, 1, len(headers))

    status_color = {
        'A Fazer': '6C757D',
        'Fazendo': '0D6EFD',
        'Concluído': '198754',
        'Ajustes': 'FFC107',
    }

    for t in tasks:
        status_label = STATUS_LABEL.get(t.status, t.status)
        ws3.append([
            t.id, t.title,
            (t.description or '').replace('\n', ' '),
            status_label,
            t.priority.capitalize(),
            t.owner.username if t.owner else '—',
            t.assignee.username if t.assignee else '—',
            t.created_at.strftime('%d/%m/%Y %H:%M'),
            t.completed_at.strftime('%d/%m/%Y %H:%M') if t.completed_at else '',
        ])
        last = ws3.max_row
        for col in range(1, len(headers) + 1):
            c = ws3.cell(row=last, column=col)
            c.border = border
            c.alignment = left
        c = ws3.cell(row=last, column=4)
        c.fill = PatternFill('solid', fgColor=status_color.get(status_label, 'FFFFFF'))
        c.font = Font(
            color='FFFFFF' if status_label != 'Ajustes' else '000000',
            bold=True
        )
        c.alignment = center
    autosize(ws3)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=_filename('xlsx'),
    )


# ============================================================
# PDF
# ============================================================

@reports_bp.route('/pdf')
@login_required
def export_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    )

    tasks, total, by_status, by_priority, grouped = _get_tasks_and_stats()

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4),
        leftMargin=1 * cm, rightMargin=1 * cm,
        topMargin=1 * cm, bottomMargin=1 * cm,
        title='Relatório Quadro de Tarefas',
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle', parent=styles['Title'],
        textColor=colors.HexColor('#0D6EFD'), fontSize=18
    )
    h2 = ParagraphStyle(
        'H2', parent=styles['Heading2'],
        textColor=colors.HexColor('#343A40'), fontSize=13, spaceAfter=6
    )
    small = ParagraphStyle('Small', parent=styles['Normal'], fontSize=9, textColor=colors.grey)

    story = []
    story.append(Paragraph('Relatório Quadro de Tarefas', title_style))
    story.append(Paragraph(
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} "
        f"por <b>{current_user.username}</b>", small
    ))
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph('Resumo por Status', h2))
    data = [['Status', 'Quantidade']] + [[k, str(v)] for k, v in by_status.items()]
    data.append(['TOTAL', str(total)])
    t1 = Table(data, colWidths=[6 * cm, 4 * cm])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343A40')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#F8F9FA')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8F9FA')]),
    ]))
    story.append(t1)
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph('Resumo por Prioridade', h2))
    data = [['Prioridade', 'Quantidade']] + [[k, str(v)] for k, v in by_priority.items()]
    t2 = Table(data, colWidths=[6 * cm, 4 * cm])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343A40')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
    ]))
    story.append(t2)
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph('Tarefas por Responsável', h2))
    resp = [['Responsável', 'Total', 'A Fazer', 'Fazendo', 'Concluído', 'Ajustes']]
    for name, d in grouped.items():
        resp.append([name, str(d['total']), str(d['todo']),
                     str(d['doing']), str(d['done']), str(d['adjust'])])
    t3 = Table(resp)
    t3.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343A40')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
    ]))
    story.append(t3)
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph('Todas as Tarefas', h2))
    td = [['#', 'Título', 'Status', 'Prioridade', 'Criado por', 'Responsável', 'Criada em']]
    for t in tasks:
        td.append([
            str(t.id),
            t.title[:60] + ('…' if len(t.title) > 60 else ''),
            STATUS_LABEL.get(t.status, t.status),
            t.priority.capitalize(),
            t.owner.username if t.owner else '—',
            t.assignee.username if t.assignee else '—',
            t.created_at.strftime('%d/%m/%Y %H:%M'),
        ])
    t4 = Table(td, repeatRows=1)
    t4.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343A40')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t4)

    doc.build(story)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=_filename('pdf'),
    )